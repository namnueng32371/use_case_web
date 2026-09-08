import glob
import openpyxl

path = glob.glob('use_caseทั้งหมด/AI Transformation*.xlsx')[0]
wb = openpyxl.load_workbook(path, data_only=True)

keywords = ['control', 'ควบคุม', 'kqi', 'qi ', 'ชี้วัด', 'ตัวชี้วัด', 'จุดควบคุม']

with open('C:/Users/Jengza/AppData/Local/Temp/claude/control_kqi_search.txt', 'w', encoding='utf-8') as f:
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        f.write(f'=== SHEET: {sheetname} ===\n')
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                s = str(v)
                low = s.lower()
                for kw in keywords:
                    if kw in low:
                        f.write(f'{cell.coordinate}: {s[:400]}\n')
                        break
print('done')
